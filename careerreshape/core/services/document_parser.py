"""Safe extraction of plain text from uploaded resume/profile files.

Two things the original implementation skipped: an upfront size check on
the raw bytes (a huge upload shouldn't even be handed to a parser), and a
guard against zip-bomb-style DOCX files (a .docx is a zip container --
its *compressed* upload size can be tiny while decompressing to
gigabytes).
"""
from __future__ import annotations

import io
import zipfile

import docx
from pypdf import PdfReader

from careerreshape.config import Settings
from careerreshape.core.exceptions import (
    FileTooLargeError,
    MalformedDocumentError,
    UnsupportedFileTypeError,
)
from careerreshape.core.models import ParsedDocument


def parse_document(file_bytes: bytes, filename: str, settings: Settings) -> ParsedDocument:
    if len(file_bytes) > settings.max_upload_bytes:
        raise FileTooLargeError(
            f"{filename} is {len(file_bytes)} bytes, exceeding the "
            f"{settings.max_upload_bytes} byte limit"
        )

    name = filename.lower()
    if name.endswith(".pdf"):
        text = _parse_pdf(file_bytes)
    elif name.endswith(".docx"):
        text = _parse_docx(file_bytes, settings)
    elif name.endswith(".txt"):
        text = file_bytes.decode("utf-8", errors="ignore")
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type for {filename!r}. Please upload a PDF, DOCX, or TXT file."
        )

    text = text.strip()
    truncated = len(text) > settings.max_input_chars
    if truncated:
        text = text[: settings.max_input_chars]

    return ParsedDocument(
        text=text,
        source_filename=filename,
        char_count=len(text),
        truncated=truncated,
    )


def _parse_pdf(file_bytes: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise MalformedDocumentError(f"Could not read PDF: {exc}") from exc


def _parse_docx(file_bytes: bytes, settings: Settings) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            total_uncompressed = sum(info.file_size for info in archive.infolist())
            if total_uncompressed > settings.max_decompressed_docx_bytes:
                raise MalformedDocumentError(
                    "DOCX file expands far beyond a normal document size and was rejected"
                )
        document = docx.Document(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in document.paragraphs)
    except zipfile.BadZipFile as exc:
        raise MalformedDocumentError(f"Not a valid DOCX file: {exc}") from exc
    except MalformedDocumentError:
        raise
    except Exception as exc:
        raise MalformedDocumentError(f"Could not read DOCX: {exc}") from exc
