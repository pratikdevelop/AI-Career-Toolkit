"""Builds a formatted .docx from optimized resume plain text."""
from __future__ import annotations

import io

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def build_resume_docx(resume_text: str) -> io.BytesIO:
    document = docx.Document()

    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    first_line_used_as_title = False
    for line in resume_text.split("\n"):
        stripped = line.strip()
        if not stripped:
            document.add_paragraph("")
            continue

        if not first_line_used_as_title:
            heading = document.add_heading(stripped, level=1)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
            first_line_used_as_title = True
            continue

        if stripped.isupper() and len(stripped) < 40:
            document.add_heading(stripped.title(), level=2)
            continue

        if stripped.startswith(("-", "\u2022", "*")):
            document.add_paragraph(stripped.lstrip("-\u2022* ").strip(), style="List Bullet")
            continue

        document.add_paragraph(stripped)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer
